from pathlib import Path
import pandas as pd
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter


# Define a pasta principal do projeto
BASE_DIR = Path(__file__).resolve().parent.parent

# Arquivo de entrada
arquivo = BASE_DIR / "dados" / "controle_vendas.xlsx"

# Pasta de saída
pasta_relatorios = BASE_DIR / "relatorios"

# Cria a pasta de relatórios caso ela não exista
pasta_relatorios.mkdir(exist_ok=True)


def carregar_dados(arquivo):
    """Lê os dados da planilha Excel."""

    try:
        return pd.read_excel(arquivo)

    except FileNotFoundError:
        print("\nERRO: Arquivo de vendas não encontrado.")
        print("Verifique se o arquivo está na pasta dados.")
        return None

df = carregar_dados(arquivo)

if df is None:
    exit()

print("Dados da planilha:")
print(df)

print("\nTotal de registros:", len(df))


# Calcula o valor total de cada venda
def calcular_total(df):
    """Calcula o valor total de cada venda."""
    df["Total"] = df["Quantidade"] * df["Valor"]
    return df


df = calcular_total(df)

print("\nValor total por venda:")
print(df[["Vendedor", "Produto", "Quantidade", "Valor", "Total"]])


# Identifica a maior venda
def identificar_maior_venda(df):
    """Identifica a venda de maior valor."""
    return df.loc[df["Total"].idxmax()]


maior_venda = identificar_maior_venda(df)

print("\nMaior venda:")
print(f"Vendedor: {maior_venda['Vendedor']}")
print(f"Produto: {maior_venda['Produto']}")
print(f"Valor: R$ {maior_venda['Total']:.2f}")


# Calcula os indicadores
def calcular_indicadores(df):
    """Calcula os principais indicadores das vendas."""
    faturamento_total = df["Total"].sum()
    media_vendas = df["Total"].mean()
    quantidade_total = df["Quantidade"].sum()

    return faturamento_total, media_vendas, quantidade_total


faturamento_total, media_vendas, quantidade_total = calcular_indicadores(df)


# Cria o resumo
def criar_resumo(faturamento_total, media_vendas, quantidade_total, maior_venda):
    """Cria o DataFrame com os principais indicadores."""
    return pd.DataFrame({
        "Indicador": [
            "Faturamento total",
            "Média por venda",
            "Quantidade total de produtos",
            "Maior venda"
        ],
        "Valor": [
            faturamento_total,
            media_vendas,
            quantidade_total,
            maior_venda["Total"]
        ]
    })


resumo = criar_resumo(
    faturamento_total,
    media_vendas,
    quantidade_total,
    maior_venda
)


# Ordena as vendas da maior para a menor
relatorio = df.sort_values("Total", ascending=False)


# Classifica as vendas
def classificar_vendas(df, valor_minimo=2000):
    """Define o status das vendas e identifica os alertas."""

    df["Status"] = df["Total"].apply(
        lambda valor: "Normal" if valor >= valor_minimo else "Atenção"
    )

    alertas = df[df["Total"] < valor_minimo]

    return df, alertas


df, alertas = classificar_vendas(df)

print("\nStatus das vendas:")
print(df[["Vendedor", "Total", "Status"]])


# Gera o relatório Excel
def gerar_relatorio(resumo, relatorio, alertas):
    """Gera o arquivo Excel com os dados do relatório."""

    with pd.ExcelWriter(
        pasta_relatorios / "relatorio_vendas.xlsx",
        engine="openpyxl"
    ) as writer:

        # Cria a aba Resumo
        resumo.to_excel(
            writer,
            sheet_name="Resumo",
            index=False
        )

        # Cria a aba Vendas
        relatorio.to_excel(
            writer,
            sheet_name="Vendas",
            index=False
        )

        # Cria a aba Alertas
        alertas.to_excel(
            writer,
            sheet_name="Alertas",
            index=False
        )

        # Acessa o arquivo Excel
        workbook = writer.book

        # Cria um gráfico de vendas
        grafico = BarChart()

        grafico.title = "Vendas por vendedor"
        grafico.y_axis.title = "Valor da venda"
        grafico.x_axis.title = "Vendedor"

        # Seleciona os valores da coluna Total
        dados_grafico = Reference(
            workbook["Vendas"],
            min_col=5,
            min_row=1,
            max_row=len(relatorio) + 1
        )

        # Seleciona os vendedores
        categorias = Reference(
            workbook["Vendas"],
            min_col=1,
            min_row=2,
            max_row=len(relatorio) + 1
        )

        grafico.add_data(
            dados_grafico,
            titles_from_data=True
        )

        grafico.set_categories(categorias)

        workbook["Vendas"].add_chart(grafico, "G2")

        # Formatação das planilhas
        for sheet in workbook.worksheets:

            # Mantém o cabeçalho visível
            sheet.freeze_panes = "A2"

            # Formata o cabeçalho
            for cell in sheet[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")

                if sheet.title == "Alertas":
                    cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor="FFC7CE"
                    )

            # Formata valores monetários
            if sheet.title == "Resumo":
                colunas_monetarias = [2]

            elif sheet.title in ["Vendas", "Alertas"]:
                colunas_monetarias = [4, 5]

            else:
                colunas_monetarias = []

            # Formata os valores das células
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.column in colunas_monetarias:
                        if isinstance(cell.value, (int, float)):
                            cell.number_format = 'R$ #,##0.00'

            # Quantidade total de produtos não é moeda
            if sheet.title == "Resumo":
                sheet["B4"].number_format = '0'

            # Ajusta automaticamente a largura das colunas
            for column in sheet.columns:

                largura = 0

                for cell in column:

                    if cell.value is not None:
                        largura = max(
                            largura,
                            len(str(cell.value))
                        )

                letra = get_column_letter(column[0].column)

                sheet.column_dimensions[letra].width = largura + 3


# Executa a geração do relatório
gerar_relatorio(resumo, relatorio, alertas)


# Mensagens finais
print("\nRelatório gerado com sucesso!")

print("\nIndicadores:")
print(f"Faturamento total: R$ {faturamento_total:.2f}")
print(f"Média por venda: R$ {media_vendas:.2f}")
print(f"Quantidade total de produtos: {quantidade_total}")
print(f"Quantidade de alertas: {len(alertas)}")